import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SHEETS = ["Inventory", "Area scores", "Gaps", "90-day plan", "Assumptions"]


def pdf_pages(docx: Path, out: Path) -> int:
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(docx)],
                   check=True, capture_output=True, timeout=180)
    import pdfplumber
    with pdfplumber.open(out / (docx.stem + ".pdf")) as pdf:
        return len(pdf.pages)


@pytest.mark.parametrize("name", ["fintech", "nbfc"])
def test_all_four_outputs_are_created(packs, name):
    files = packs[name].files
    assert {k: v.name for k, v in files.items()} == {"workbook": "evidence_pack.xlsx", "heatmap": "heatmap.png",
                                                     "report": "evidence_pack.docx", "board_summary": "board_summary.docx"}
    assert all(f.exists() and f.stat().st_size > 5000 for f in files.values())


def test_workbook_has_the_five_sheets_and_one_row_per_use_case(packs):
    for name, n in (("fintech", 8), ("nbfc", 12)):
        wb = load_workbook(packs[name].files["workbook"])
        assert wb.sheetnames == SHEETS
        inv = wb["Inventory"]
        assert inv.max_row == n + 1
        head = [c.value for c in inv[1]]
        assert "Assessed risk tier" in head and "Business owner" in head
        tiers = {inv.cell(row=r, column=head.index("Assessed risk tier") + 1).value for r in range(2, n + 2)}
        assert tiers <= {"High", "Medium", "Low"}


def test_gaps_sheet_columns_and_blank_due_dates(packs):
    for pack in packs.values():
        ws = load_workbook(pack.files["workbook"])["Gaps"]
        head = [c.value for c in ws[1]]
        assert {"Gap", "Area", "Severity", "Owner"} <= set(head)
        due = head.index("Due date (client to set)") + 1
        assert ws.max_row == len(pack.scored.gaps) + 1
        assert all(ws.cell(row=r, column=due).value in (None, "") for r in range(2, ws.max_row + 1))
        sev = [ws.cell(row=r, column=head.index("Severity") + 1).value for r in range(2, ws.max_row + 1)]
        assert set(sev) <= {"High", "Medium", "Low"}


def test_area_scores_sheet_has_matrix_summary_and_evidence(packs):
    ws = load_workbook(packs["fintech"].files["workbook"])["Area scores"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    for a in ("Accountability", "AI inventory", "Data governance", "Lifecycle and evals", "AI security", "Human oversight",
              "Vendor AI risk", "Monitoring and audit"):
        assert a in values
    assert any(isinstance(v, str) and "to verify against the RBI FREE-AI report text before client use" in v for v in values)
    assert any(isinstance(v, str) and v.startswith("Overall readiness") for v in values)
    assert any(isinstance(v, str) and v.startswith("Evidence notes") for v in values)


def test_plan_and_assumptions_sheets_not_empty(packs):
    wb = load_workbook(packs["nbfc"].files["workbook"])
    assert wb["90-day plan"].max_row > 5 and wb["Assumptions"].max_row > 8


def test_heatmap_is_a_png_of_sensible_size(packs):
    for pack in packs.values():
        with Image.open(pack.files["heatmap"]) as im:
            assert im.format == "PNG" and im.width >= 1500 and im.height >= 800


def test_report_has_required_sections_and_disclaimer(packs):
    doc = Document(packs["fintech"].files["report"])
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for h in ("Executive summary", "Scope", "Method", "Inventory overview", "Findings by area", "Top 10 gaps", "90-day plan",
              "Limits of this report"):
        assert any(h in x for x in heads), h
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "not a formal audit" in text and "CERT-In empanelled auditors" in text
    assert sum(1 for x in heads if x[0].isdigit()) >= 8   # numbered sections 1-8 (area headings are numbered too)
    assert len(doc.inline_shapes) == 1                                        # the heatmap
    top10 = next(t for t in doc.tables if t.rows[0].cells[1].text == "Gap")
    assert 1 < len(top10.rows) <= 11


def test_board_summary_content(packs):
    doc = Document(packs["nbfc"].files["board_summary"])
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for s in ("Board summary", "Top 5 gaps", "Decisions we ask the board to take", "90-day plan at a glance", "formal audit"):
        assert s in text


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed: page-count test skipped")
def test_page_counts(packs, tmp_path):
    for pack in packs.values():
        assert 6 <= pdf_pages(pack.files["report"], tmp_path) <= 10
        assert pdf_pages(pack.files["board_summary"], tmp_path) == 2


def test_outputs_are_reproducible_from_the_same_input(cfg, tmp_path):
    from evidence_pack.pipeline import build_pack
    a = build_pack(ROOT / "samples/fintech_120.yaml", tmp_path / "a", config=cfg).scored
    b = build_pack(ROOT / "samples/fintech_120.yaml", tmp_path / "b", config=cfg).scored
    assert a.overall == b.overall and [(g.check_id, g.use_case_id, g.severity_score) for g in a.gaps] == \
        [(g.check_id, g.use_case_id, g.severity_score) for g in b.gaps]


def test_sample_sizes_match_the_brief(packs):
    assert len(packs["fintech"].scored.register.use_cases) == 8 and "120 staff" in packs["fintech"].scored.register.org.size
    assert len(packs["nbfc"].scored.register.use_cases) == 12


def test_no_real_organisations_in_samples():
    for f in ("fintech_120.yaml", "nbfc_12.csv", "nbfc_org.yaml"):
        text = (ROOT / "samples" / f).read_text().lower()
        assert "(demo" in text and "@" not in text


def test_cli_success_and_error_exit_codes(tmp_path):
    ok = subprocess.run([sys.executable, "-m", "evidence_pack", "samples/fintech_120.yaml", "--out", str(tmp_path / "o")],
                        cwd=ROOT, capture_output=True, text=True)
    assert ok.returncode == 0 and "8 use cases" in ok.stdout and (tmp_path / "o/evidence_pack.xlsx").exists()
    bad = tmp_path / "bad.csv"
    bad.write_text("ID,Name,Decision impact\nA,n,maybe\n")
    err = subprocess.run([sys.executable, "-m", "evidence_pack", str(bad), "--out", str(tmp_path / "o2")],
                         cwd=ROOT, capture_output=True, text=True)
    assert err.returncode == 2 and "decision_impact" in err.stderr and not (tmp_path / "o2").exists()
