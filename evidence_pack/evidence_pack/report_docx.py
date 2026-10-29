"""evidence_pack.docx: the full readiness report (target 6-10 pages)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx.shared import Cm

from . import narrative as nv
from .docx_util import MUTED, add_table, bullet, footer, new_document, para
from .models import Scored
from .palette import SCORE_FILL, SCORE_TEXT, SEVERITY_FILL, SEVERITY_TEXT, level
from .scoring import Config

MAX_PLAN_ROWS = 14


def _ids(ids: list[str], n: int = 4) -> str:
    return ", ".join(ids[:n]) + (f" +{len(ids) - n} more" if len(ids) > n else "")


def write_report(sc: Scored, cfg: Config, heatmap_png: Path, path: Path) -> Path:
    org, ucs = sc.register.org, sc.register.use_cases
    verify = cfg.areas_doc["verify_note"]
    doc = new_document(10, 2.0)
    footer(doc, f"{org.name}: AI readiness report (readiness advice, not a formal audit)")

    # ---------------- title + executive summary
    para(doc, "AI readiness report", bold=True, size=24, after=0)
    para(doc, "Bank-ready AI evidence pack", size=13, color=MUTED, after=2)
    para(doc, f"{org.name}  |  {org.sector}  |  {org.size}".strip(" |"), size=10.5, after=0)
    para(doc, f"Assessment date: {(org.as_of or date.today()).strftime('%d %B %Y')}"
              + (f"  |  Prepared by: {org.prepared_by}" if org.prepared_by else ""), size=9.5, color=MUTED, after=6)
    doc.add_heading("1. Executive summary", 1)
    for s in nv.exec_summary(sc, cfg):
        bullet(doc, s)
    tc = nv.tier_counts(sc)
    para(doc, "", after=2)
    add_table(doc, ["Overall score", "Use cases", "High / Medium / Low risk", "Gaps (High)", "Plan actions"],
              [[f"{sc.overall:.1f} / 4", str(len(ucs)), f"{tc['High']} / {tc['Medium']} / {tc['Low']}",
                f"{len(sc.gaps)} ({sum(g.severity == 'High' for g in sc.gaps)})", str(len(sc.plan))]],
              [3.2, 2.6, 4.4, 3.4, 3.4], size=10, center_cols={0, 1, 2, 3, 4})
    para(doc, "", after=2)
    rows = [[a["name"], f"{sc.area_avg[a['key']]:.1f}", nv.label(cfg, sc.area_avg[a["key"]]),
             ", ".join(u.id for u in ucs if sc.scores[u.id][a["key"]].score <= 2) or "none"] for a in cfg.areas]
    fills = {(i, 1): (SCORE_FILL[level(sc.area_avg[a["key"]])], SCORE_TEXT[level(sc.area_avg[a["key"]])])
             for i, a in enumerate(cfg.areas)}
    add_table(doc, ["Area", "Score", "Level", "Use cases scoring 1 or 2"], rows, [4.6, 1.6, 3.6, 7.2], fills=fills,
              center_cols={1})

    # ---------------- scope and method
    doc.add_page_break()
    doc.add_heading("2. Scope", 1)
    para(doc, f"In scope: the {len(ucs)} AI use cases in the register supplied by {org.name} "
              f"({', '.join(u.id for u in ucs)}), and four organisation-level controls: board-approved AI policy, AI "
              "governance owner, inventory review cycle and risk-based audit coverage.")
    para(doc, "Out of scope: testing of any model, system or vendor; review of source code, infrastructure or contracts; "
              "interviews; and any formal security audit. Every answer is taken as the organisation stated it.")
    doc.add_heading("3. Method", 1)
    para(doc, "Each use case is scored 1 to 4 on eight areas. A score is the share of applicable, weighted checks the use case "
              "earns, banded as follows: " + ", ".join(
                  f"{lv} {cfg.areas_doc['score_labels'][lv].lower()} (from {cfg.scoring['bands'][lv]:.0%})" for lv in (1, 2, 3, 4))
         + ". A blank answer scores as not evidenced. Checks that do not apply to a use case are left out of its score.")
    add_table(doc, ["#", "Area", "What it covers", "Reference (" + verify + ")"],
              [[str(a["id"]), a["name"], a["covers"], "; ".join(a["references"])] for a in cfg.areas],
              [0.8, 3.2, 7.0, 6.0], size=8.5)
    para(doc, "", after=2)
    para(doc, "Risk tiers are assigned by the rules below (the highest matching tier wins). They set how strict the checks are "
              "(for example, how recent a validation must be) and how severe a gap is.", after=3)
    add_table(doc, ["Rule", "Tier", "Condition"],
              [[r["id"], r["tier"], r["description"]] for r in cfg.tiering["rules"]] +
              [["default", cfg.tiering["default_tier"], "No rule matched"]], [1.6, 2.0, 13.4], size=8.5)
    para(doc, "", after=2)
    t = cfg.scoring["severity_thresholds"]
    para(doc, f"Gap severity = check weight x shortfall x tier factor (High {cfg.tiering['tier_factor']['High']}, "
              f"Medium {cfg.tiering['tier_factor']['Medium']}, Low {cfg.tiering['tier_factor']['Low']}). "
              f"A result of {t['High']:g} or more is High, {t['Medium']:g} or more is Medium, otherwise Low. "
              "Area scores across use cases are tier-weighted means.", size=9, color=MUTED)

    # ---------------- inventory overview
    doc.add_page_break()
    doc.add_heading("4. Inventory overview", 1)
    para(doc, f"{len(ucs)} use cases: {tc['High']} High, {tc['Medium']} Medium, {tc['Low']} Low. "
              f"{sum(1 for u in ucs if u.uses_genai)} use generative AI, {sum(1 for u in ucs if u.third_party)} rely on a "
              f"third-party model or vendor, and {sum(1 for u in ucs if u.personal_effective)} use personal data.")
    doc.add_picture(str(heatmap_png), width=Cm(17.0))
    inv_rows, fills = [], {}
    for i, u in enumerate(ucs):
        avg = sum(a.score for a in sc.scores[u.id].values()) / 8
        diff = f" (client: {u.client_risk_tier})" if u.client_risk_tier and u.client_risk_tier != sc.tiers[u.id].tier else ""
        inv_rows.append([u.id, u.name, sc.tiers[u.id].tier + diff, f"{avg:.1f}"])
        fills[(i, 2)] = (SEVERITY_FILL[sc.tiers[u.id].tier], SEVERITY_TEXT[sc.tiers[u.id].tier])
        fills[(i, 3)] = (SCORE_FILL[level(avg)], SCORE_TEXT[level(avg)])
    para(doc, "", after=2)
    add_table(doc, ["ID", "Use case", "Assessed tier", "Average"], inv_rows, [1.8, 8.4, 4.6, 2.2], size=8.5, fills=fills,
              center_cols={3})
    para(doc, "Where the assessed tier differs from the client's own, the client's tier is shown in brackets.",
         italic=True, size=8.5, color=MUTED)

    # ---------------- area findings
    doc.add_page_break()
    doc.add_heading("5. Findings by area", 1)
    extra = cfg.areas_doc.get("additional_references", {})
    for a in cfg.areas:
        score = sc.area_avg[a["key"]]
        doc.add_heading(f"{a['id']}. {a['name']}: {score:.1f} / 4 ({nv.label(cfg, score)})", 2)
        refs = "; ".join(a["references"] + extra.get(a["id"], []))
        para(doc, f"{a['covers']}. Reference: {refs} ({verify}).", italic=True, size=9, color=MUTED, after=2)
        f = nv.area_findings(sc, cfg, a["key"])
        for r in f["checks"]:
            bullet(doc, nv.check_sentence(r), size=9.5)
        para(doc, "Use-case scores: " + ", ".join(f"{k} {v}" for k, v in f["scores"].items()) + ".", size=9, after=3)

    # ---------------- top 10 gaps
    doc.add_page_break()
    doc.add_heading("6. Top 10 gaps", 1)
    para(doc, "Gaps that repeat across use cases are grouped into one theme and ranked by their combined severity score. "
              "All individual gaps, with owners and blank due dates, are in the Gaps sheet of the workbook.", size=9.5)
    top = sc.themes[:10]
    rows = [[str(t["rank"]), t["title"], cfg.area_by_key[t["area"]]["name"],
             "Organisation" if t["org_level"] else f"{t['count']}: {_ids(t['use_cases'])}", t["severity"]] for t in top]
    fills = {(i, 4): (SEVERITY_FILL[t["severity"]], SEVERITY_TEXT[t["severity"]]) for i, t in enumerate(top)}
    add_table(doc, ["#", "Gap", "Area", "Use cases", "Severity"], rows, [0.9, 7.4, 3.2, 4.2, 1.9], size=9, fills=fills,
              center_cols={0})

    # ---------------- 90-day plan
    doc.add_heading("7. 90-day plan", 1)
    para(doc, "Actions are phased by severity: High-severity themes in the first 30 days, Medium in days 31-60, Low in days "
              "61-90. Owners are suggestions for the client to confirm. Effort: S = days, M = weeks, L = a month or more.",
         size=9.5)
    plan = sc.plan[:MAX_PLAN_ROWS]
    rows = [[p["phase"], p["action"], p["owner"], p["effort"]] for p in plan]
    add_table(doc, ["Phase", "Action", "Suggested owner", "Effort"], rows, [2.1, 9.4, 4.4, 1.3], size=8.5, center_cols={3})
    if len(sc.plan) > MAX_PLAN_ROWS:
        para(doc, f"{len(sc.plan) - MAX_PLAN_ROWS} further lower-priority actions are listed in the 90-day plan sheet.",
             italic=True, size=8.5, color=MUTED)

    # ---------------- disclaimer
    doc.add_heading("8. Limits of this report and disclaimer", 1)
    para(doc, nv.DISCLAIMER, size=9.5)
    para(doc, "Assumptions used in scoring:", bold=True, size=9.5, after=2)
    for a in sc.assumptions[:8]:
        bullet(doc, a, size=8.5)
    para(doc, "The full list of assumptions is in the Assumptions sheet of the workbook.", italic=True, size=8.5, color=MUTED)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
