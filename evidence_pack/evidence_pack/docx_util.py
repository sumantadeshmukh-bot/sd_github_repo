"""Small python-docx helpers shared by the report and the board summary."""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x52, 0x60, 0x6D)
HEAD_HEX = "1F2933"


def new_document(base_pt: float = 10, margin_cm: float = 2.0) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
    sec.left_margin = sec.right_margin = Cm(margin_cm)
    sec.top_margin, sec.bottom_margin = Cm(margin_cm), Cm(margin_cm * 0.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Calibri", Pt(base_pt), INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before in (("Heading 1", 16, 10), ("Heading 2", 12.5, 8), ("Heading 3", 11, 6)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = "Calibri", Pt(size), True, INK
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(3)
        st.paragraph_format.keep_with_next = True
    return doc


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell(cell, text: str, bold: bool = False, color: RGBColor | None = None, size: float = 9,
             align: str = "left") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    run = p.add_run(str(text))
    run.bold, run.font.size = bold, Pt(size)
    if color is not None:
        run.font.color.rgb = color


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _no_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float], size: float = 9,
              fills: dict[tuple[int, int], tuple[str, str]] | None = None, center_cols: set[int] | None = None):
    """Table with a dark header row. ``fills[(row, col)] = (fill_hex, text_hex)`` colours individual body cells."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, bold=True, color=RGBColor(255, 255, 255), size=size)
        shade(t.rows[0].cells[i], HEAD_HEX)
    _repeat_header(t.rows[0])
    for r, row in enumerate(rows):
        cells = t.add_row().cells
        _no_split(t.rows[-1])
        for c, v in enumerate(row):
            fill = (fills or {}).get((r, c))
            set_cell(cells[c], v, bold=bool(fill), color=RGBColor.from_string(fill[1]) if fill else None, size=size,
                     align="center" if c in (center_cols or set()) else "left")
            if fill:
                shade(cells[c], fill[0])
    for i, w in enumerate(widths_cm):  # grid widths too: LibreOffice and Word both honour tblGrid
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    return t


def bullet(doc: Document, text: str, bold_prefix: str = "", size: float | None = None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        if size:
            r.font.size = Pt(size)
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    return p


def para(doc: Document, text: str, bold: bool = False, italic: bool = False, size: float | None = None,
         color: RGBColor | None = None, after: float | None = None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def footer(doc: Document, text: str) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text + "  |  Page ")
    r.font.size, r.font.color.rgb = Pt(8), MUTED
    run = p.add_run()
    run.font.size = Pt(8)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)
