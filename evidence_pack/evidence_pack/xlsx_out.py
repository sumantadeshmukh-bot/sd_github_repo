"""evidence_pack.xlsx: Inventory, Area scores, Gaps, 90-day plan, Assumptions."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Scored
from .palette import SCORE_FILL, SCORE_TEXT, SEVERITY_FILL, SEVERITY_TEXT, level
from .scoring import Config

HEAD_FILL = PatternFill("solid", fgColor="1F2933")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _header(ws, row: int, headers: list[str]) -> None:
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font, cell.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center")


def _widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, row: int, text: str, note: str = "") -> None:
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, size=12)
    if note:
        ws.cell(row=row + 1, column=1, value=note).font = Font(italic=True, color="52606D")


def _score_cell(cell, v: float) -> None:
    lv = level(v)
    cell.fill, cell.font = _fill(SCORE_FILL[lv]), Font(bold=True, color=SCORE_TEXT[lv])
    cell.alignment = Alignment(horizontal="center", vertical="top")


def write_workbook(sc: Scored, cfg: Config, path: Path) -> Path:
    wb = Workbook()
    ucs, org = sc.register.use_cases, sc.register.org
    yn = lambda v: "" if v is None else ("yes" if v else "no")  # noqa: E731

    # ---------------- Inventory
    ws = wb.active
    ws.title = "Inventory"
    heads = ["ID", "Use case", "Business owner", "Purpose", "Model / vendor", "Data categories", "Personal data",
             "Sensitive data", "Users affected", "Decision impact", "Human in the loop", "Hosting / data location",
             "Monitoring", "Last validation", "Incidents (12m)", "Client's own tier", "Assessed risk tier",
             "Why this tier", "Tier differs from client's?", "Average area score"]
    _header(ws, 1, heads)
    for r, u in enumerate(ucs, 2):
        t = sc.tiers[u.id]
        avg = sum(a.score for a in sc.scores[u.id].values()) / len(cfg.areas)
        row = [u.id, u.name, u.business_owner, u.purpose, u.model_vendor, "; ".join(u.data_categories),
               yn(u.personal_data) or f"{yn(u.personal_effective)} (inferred)", yn(u.sensitive_data) or
               f"{yn(u.sensitive_effective)} (inferred)", ", ".join(u.users_affected), u.decision_impact,
               yn(u.human_in_loop), u.hosting_location, u.monitoring, u.last_validation_date.isoformat() if
               u.last_validation_date else "", "" if u.incidents_12m is None else u.incidents_12m,
               u.client_risk_tier, t.tier, "; ".join(f"{rid}: {d}" for rid, d in t.matched),
               "" if not u.client_risk_tier else ("yes" if u.client_risk_tier != t.tier else "no"), round(avg, 1)]
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
        sev = {"High": "High", "Medium": "Medium", "Low": "Low"}[t.tier]
        tc = ws.cell(row=r, column=17)
        tc.fill, tc.font = _fill(SEVERITY_FILL[sev]), Font(bold=True, color=SEVERITY_TEXT[sev])
        _score_cell(ws.cell(row=r, column=20), avg)
    _widths(ws, [8, 28, 30, 38, 26, 32, 14, 14, 16, 14, 12, 26, 12, 14, 11, 12, 12, 52, 14, 12])
    ws.freeze_panes, ws.auto_filter.ref = "C2", f"A1:{get_column_letter(len(heads))}{len(ucs) + 1}"

    # ---------------- Area scores
    ws = wb.create_sheet("Area scores")
    _title(ws, 1, f"{org.name}: scores by use case (1 = {cfg.areas_doc['score_labels'][1].lower()}, "
                  f"4 = {cfg.areas_doc['score_labels'][4].lower()})",
           "Each score is the share of applicable weighted checks earned, banded 1-4. Evidence notes are in the third table below.")
    _header(ws, 3, ["ID", "Use case", "Tier"] + [a["name"] for a in cfg.areas] + ["Average"])
    for r, u in enumerate(ucs, 4):
        ws.cell(row=r, column=1, value=u.id)
        ws.cell(row=r, column=2, value=u.name)
        ws.cell(row=r, column=3, value=sc.tiers[u.id].tier)
        for c, a in enumerate(cfg.areas, 4):
            cell = ws.cell(row=r, column=c, value=sc.scores[u.id][a["key"]].score)
            _score_cell(cell, cell.value)
        _score_cell(ws.cell(row=r, column=12, value=round(sum(x.score for x in sc.scores[u.id].values()) / 8, 1)),
                    sum(x.score for x in sc.scores[u.id].values()) / 8)
    r0 = len(ucs) + 6
    _title(ws, r0, "Area summary (tier-weighted: High x3, Medium x2, Low x1)")
    _header(ws, r0 + 1, ["#", "Area", "What it covers", "Score (1-4)", "Level", "Use cases scoring 1 or 2",
                         "Reference mapping (" + cfg.areas_doc["verify_note"] + ")"])
    for i, a in enumerate(cfg.areas):
        r = r0 + 2 + i
        score = sc.area_avg[a["key"]]
        weak = [u.id for u in ucs if sc.scores[u.id][a["key"]].score <= 2]
        extra = cfg.areas_doc.get("additional_references", {}).get(a["id"], [])
        vals = [a["id"], a["name"], a["covers"], score, cfg.areas_doc["score_labels"][level(score)],
                ", ".join(weak) or "none", "; ".join(a["references"] + extra) + " - " + cfg.areas_doc["verify_note"]]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
        _score_cell(ws.cell(row=r, column=4), score)
    r1 = r0 + 12
    ws.cell(row=r1 - 1, column=2, value=f"Overall readiness: {sc.overall} / 4 ({cfg.areas_doc['score_labels'][level(sc.overall)]})").font = Font(bold=True)
    _title(ws, r1, "Evidence notes (one line per check)")
    _header(ws, r1 + 1, ["Use case", "Area", "Area score", "Check", "Result", "Evidence note"])
    r = r1 + 2
    label = {"pass": "Pass", "partial": "Partial", "fail": "Fail", "unknown": "Not evidenced", "na": "Not applicable"}
    for u in ucs:
        for a in cfg.areas:
            ar = sc.scores[u.id][a["key"]]
            for cr in ar.results:
                for c, v in enumerate([u.id, a["name"], ar.score, f"{cr.check_id} {cfg.checks[cr.check_id]['title']}",
                                       label[cr.result], cr.note], 1):
                    ws.cell(row=r, column=c, value=v).alignment = WRAP
                r += 1
    _widths(ws, [10, 30, 34, 16, 16, 16, 18, 16, 16, 16, 16, 10])
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["F"].width = 40

    # ---------------- Gaps
    ws = wb.create_sheet("Gaps")
    heads = ["#", "Gap", "Area", "Use case", "Severity", "Severity score", "Owner", "Evidence", "Result", "Risk tier",
             "Check", "Due date (client to set)"]
    _header(ws, 1, heads)
    for r, g in enumerate(sc.gaps, 2):
        vals = [r - 1, g.gap, cfg.area_by_key[g.area]["name"], g.use_case_id, g.severity, g.severity_score, g.owner,
                g.detail, label[g.result], g.tier, g.check_id, g.due_date]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
        sc_cell = ws.cell(row=r, column=5)
        sc_cell.fill, sc_cell.font = _fill(SEVERITY_FILL[g.severity]), Font(bold=True, color=SEVERITY_TEXT[g.severity])
    _widths(ws, [5, 48, 20, 10, 10, 10, 34, 56, 14, 10, 9, 18])
    ws.freeze_panes, ws.auto_filter.ref = "C2", f"A1:{get_column_letter(len(heads))}{len(sc.gaps) + 1}"

    # ---------------- 90-day plan
    ws = wb.create_sheet("90-day plan")
    heads = ["Phase", "Action", "Area", "Addresses", "Priority", "Suggested owner", "Effort (S/M/L)", "Done when"]
    _header(ws, 1, heads)
    for r, p in enumerate(sc.plan, 2):
        for c, v in enumerate([p["phase"], p["action"], p["area"], p["addresses"], p["priority"], p["owner"],
                               p["effort"], p["done_when"]], 1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
        cell = ws.cell(row=r, column=5)
        cell.fill, cell.font = _fill(SEVERITY_FILL[p["priority"]]), Font(bold=True, color=SEVERITY_TEXT[p["priority"]])
    _widths(ws, [12, 60, 20, 24, 10, 36, 10, 40])
    ws.freeze_panes = "B2"

    # ---------------- Assumptions
    ws = wb.create_sheet("Assumptions")
    _header(ws, 1, ["#", "Assumption or limitation"])
    for r, a in enumerate(sc.assumptions, 2):
        ws.cell(row=r, column=1, value=r - 1)
        ws.cell(row=r, column=2, value=a).alignment = WRAP
    _widths(ws, [5, 130])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
